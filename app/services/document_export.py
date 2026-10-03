"""
Academic Document Exporter Service
Generates publication- and university-grade Microsoft Word (.docx) files
from generated academic assignments.
"""

import io
import re
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def add_page_number(run):
    """Inserts a dynamic PAGE number field into a run in docx"""
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def parse_inline_markdown(paragraph, text: str, font_name: str = "Times New Roman", font_size_pt: float = 12):
    """
    Parses inline bold (**bold**), italic (*italic*), and links ([title](url))
    into appropriately styled Docx Runs.
    """
    # Tokenize by bold and italic markdown tags
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|\[.*?\]\(.*?\))', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("[") and "](" in token and token.endswith(")"):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                link_text, url = m.groups()
                run = paragraph.add_run(f"{link_text} ({url})")
                run.font.color.rgb = RGBColor(0, 51, 102)
                run.underline = True
            else:
                run = paragraph.add_run(token)
        else:
            run = paragraph.add_run(token)

        run.font.name = font_name
        run.font.size = Pt(font_size_pt)

class AcademicDocxExporter:
    """
    Converts academic markdown documents into formatted Word documents
    adhering to university submission standards (Times New Roman 12pt, 1.5 line spacing, 1-inch margins).
    """

    def generate_docx(
        self,
        title: str,
        content_markdown: str,
        student_name: str = "Student Submission",
        course_name: str = "Academic Coursework",
        institution: str = "University Submission",
        date_str: str = None
    ) -> io.BytesIO:
        doc = Document()
        
        # 1. Page Margins: Standard 1-inch (72pt) all around
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)
            section.header_distance = Inches(0.5)
            section.footer_distance = Inches(0.5)

            # Header & Footer setup
            header = section.header
            header_p = header.paragraphs[0]
            header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hrun = header_p.add_run(f"{title[:40]} | Page ")
            hrun.font.name = "Times New Roman"
            hrun.font.size = Pt(9)
            hrun.font.color.rgb = RGBColor(128, 128, 128)
            add_page_number(hrun)

        if not date_str:
            date_str = datetime.utcnow().strftime("%B %d, %Y")

        word_count = len(content_markdown.split())

        # -------------------------------------------------------------
        # 2. Cover / Title Page
        # -------------------------------------------------------------
        # Vertical spacing
        for _ in range(3):
            doc.add_paragraph()

        # Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(title)
        title_run.font.name = "Times New Roman"
        title_run.font.size = Pt(22)
        title_run.bold = True
        title_run.font.color.rgb = RGBColor(20, 24, 33)

        doc.add_paragraph() # spacing

        # Subtitle / Course
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(course_name)
        sub_run.font.name = "Times New Roman"
        sub_run.font.size = Pt(14)
        sub_run.italic = True
        sub_run.font.color.rgb = RGBColor(90, 100, 110)

        for _ in range(5):
            doc.add_paragraph()

        # Metadata Box
        meta_p = doc.add_paragraph()
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_p.paragraph_format.line_spacing = 1.3
        
        details = [
            f"Author: {student_name}",
            f"Institution: {institution}",
            f"Date of Submission: {date_str}",
            f"Word Count: {word_count:,} words"
        ]
        for line in details:
            r = meta_p.add_run(line + "\n")
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(60, 60, 60)

        # Page Break after Cover Page
        doc.add_page_break()

        # -------------------------------------------------------------
        # 3. Main Content Processing
        # -------------------------------------------------------------
        lines = content_markdown.splitlines()
        is_references = False

        for raw_line in lines:
            line = raw_line.strip()
            if not line:
                continue

            # Heading 1 (# Heading)
            if line.startswith("# ") and not line.startswith("## "):
                heading_text = line[2:].strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(heading_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(16)
                run.bold = True
                run.font.color.rgb = RGBColor(20, 24, 33)

                if heading_text.lower() in ["references", "bibliography", "works cited", "reference list"]:
                    is_references = True
                else:
                    is_references = False
                continue

            # Heading 2 (## Heading)
            if line.startswith("## ") and not line.startswith("### "):
                heading_text = line[3:].strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(heading_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(13)
                run.bold = True
                run.font.color.rgb = RGBColor(40, 50, 60)
                continue

            # Heading 3 (### Heading)
            if line.startswith("### "):
                heading_text = line[4:].strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(heading_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.bold = True
                run.italic = True
                run.font.color.rgb = RGBColor(50, 60, 70)
                continue

            # Bullet points (* or -)
            if line.startswith("* ") or line.startswith("- "):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.25
                parse_inline_markdown(p, line[2:].strip(), font_name="Times New Roman", font_size_pt=12)
                continue

            # Numbered list item (1. item)
            num_match = re.match(r'^(\d+)\.\s+(.*)', line)
            if num_match:
                p = doc.add_paragraph(style='List Number')
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.25
                parse_inline_markdown(p, num_match.group(2).strip(), font_name="Times New Roman", font_size_pt=12)
                continue

            # Standard Paragraph or Reference Entry
            p = doc.add_paragraph()
            if is_references:
                # Academic hanging indent for References (0.5 inch / 36pt hanging)
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
            else:
                # Body paragraph: 1.5 line spacing, 6pt space after, justified
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.5
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            parse_inline_markdown(p, line, font_name="Times New Roman", font_size_pt=12)

        # Save to in-memory bytes buffer
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return file_stream

docx_exporter = AcademicDocxExporter()

def get_docx_exporter() -> AcademicDocxExporter:
    return docx_exporter
