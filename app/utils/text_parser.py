import io
import base64
from fastapi import UploadFile, HTTPException
import logging

logger = logging.getLogger("text_parser")

# Minimal imports to avoid heavy deps if not present
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


async def _gemini_ocr_pdf(pdf_bytes: bytes) -> str:
    """Use Gemini's multimodal vision to OCR a scanned/image-based PDF."""
    try:
        import google.generativeai as genai
        from app.config import get_settings

        settings = get_settings()
        # Explicitly configure with API key — text_parser is a standalone module
        # and may be used before ai_service.py has been imported/configured.
        genai.configure(api_key=settings.GEMINI_API_KEY)
        model = genai.GenerativeModel(settings.GEMINI_MODEL_NAME)

        pdf_b64 = base64.standard_b64encode(pdf_bytes).decode("utf-8")

        response = model.generate_content([
            {
                "inline_data": {
                    "mime_type": "application/pdf",
                    "data": pdf_b64,
                }
            },
            "Please extract and transcribe ALL text content from this document exactly as it appears. "
            "Preserve headings, paragraphs, and lists. Output only the extracted text, nothing else."
        ])
        extracted = response.text or ""
        logger.info(f"Gemini OCR returned {len(extracted)} characters")
        return extracted
    except Exception as e:
        logger.error(f"Gemini OCR fallback failed with error: {type(e).__name__}: {str(e)}")
        return ""


async def extract_text_from_file(file: UploadFile) -> str:
    filename = file.filename.lower()
    content = await file.read()
    file_stream = io.BytesIO(content)
    
    text = ""
    
    try:
        if filename.endswith(".pdf"):
            if not PdfReader:
                raise HTTPException(status_code=500, detail="PDF support not installed")
            reader = PdfReader(file_stream)
            if len(reader.pages) > 50:
                raise HTTPException(status_code=400, detail="PDF exceeds the 50-page maximum limit.")
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
            
            # ── Fallback: if pypdf returned nothing, the PDF is image/scan-based ──
            if not text.strip():
                logger.info(f"pypdf returned empty text for {filename} — using Gemini OCR fallback")
                text = await _gemini_ocr_pdf(content)
                if not text.strip():
                    raise HTTPException(
                        status_code=400,
                        detail="Could not extract text from this PDF. It may be a fully image-based scan that could not be read. Try converting it to a text-based PDF or copy the content into a .txt file."
                    )
                
        elif filename.endswith(".docx"):
            if not Document:
                raise HTTPException(status_code=500, detail="DOCX support not installed")
            doc = Document(file_stream)
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_texts = []
                    seen_cells = set()
                    for cell in row.cells:
                        if cell not in seen_cells:
                            seen_cells.add(cell)
                            c_text = cell.text.strip()
                            if c_text:
                                row_texts.append(c_text)
                    if row_texts:
                        text += " | ".join(row_texts) + "\n"

        elif filename.endswith(".pptx"):
            if not Presentation:
                raise HTTPException(status_code=500, detail="PPTX support not installed")
            prs = Presentation(file_stream)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"

        elif filename.endswith(".xlsx") or filename.endswith(".xls") or filename.endswith(".xlsm"):
            if not openpyxl:
                raise HTTPException(status_code=500, detail="Excel support not installed")
            try:
                wb = openpyxl.load_workbook(file_stream, data_only=True)
                texts = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        texts.append(" ".join([str(cell) for cell in row if cell is not None]))
                return "\n".join(texts)
            except Exception as e:
                if filename.endswith(".xls"):
                    raise HTTPException(status_code=400, detail="Standard .xls files are old. Please convert to .xlsx or upload as PDF.")
                raise e
            
        elif filename.endswith(".txt"):
            return content.decode("utf-8", errors="ignore")
            
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format. Please use PDF, DOCX, PPTX, XLSX, or plain TXT.")
            
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error parsing file {filename}: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to parse document: {str(e)}")
        
    # Reset cursor for other uses if needed (like upload to storage)
    await file.seek(0)
    
    return text.strip()
